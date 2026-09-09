from __future__ import annotations

import hashlib
import json
import math
import sys
from pathlib import Path
from typing import Any, Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = Path(
    r"C:\Users\戴周阳\.codex\plugins\cache\openai-primary-runtime\documents\26.826.12353\skills\documents"
)
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


OUTPUT = ROOT / "reports" / "deliverables" / "510300历史策略与全部回测结果报告_20260830.docx"

SOURCE_PATHS = {
    "goal_v2": ROOT / "reports" / "audit" / "510300_sharpe_1_2_goal_status_v2_20260830.json",
    "goal_v1": ROOT / "reports" / "audit" / "510300_sharpe_1_2_goal_status_20260830.json",
    "micro": ROOT / "reports" / "audit" / "510300_frozen_microstructure_sharpe_audit_v1.json",
    "micro_discovery": ROOT / "reports" / "discovery" / "510300_etf_microstructure_shadow_gate_discovery_v5.json",
    "registered": ROOT / "reports" / "backtest" / "registered_factor_backtests.json",
    "round1": ROOT / "reports" / "research" / "round1_factor_selection_decision.json",
    "round2": ROOT / "reports" / "research" / "round2_factor_selection_decision.json",
    "round3": ROOT / "reports" / "research" / "round3_factor_selection_decision.json",
    "round4": ROOT / "reports" / "research" / "round4_factor_selection_decision.json",
    "phase1": ROOT / "reports" / "backtest" / "phase1_five_year_backtest.json",
    "r6": ROOT / "reports" / "backtest" / "r6_model_selection.json",
    "overlay": ROOT / "reports" / "backtest" / "valuation_intraday_entry_overlay.json",
    "t_only": ROOT / "reports" / "backtest" / "t_only_robustness_v1.json",
    "cross_etf": ROOT / "reports" / "research" / "510300_cross_etf_forced_flow_binary_screen_v1.json",
    "atlas": ROOT / "reports" / "research" / "510300_mechanism_atlas_v1.json",
    "promotion": ROOT / "reports" / "research" / "510300_mechanism_promotion_gate_adjudication_v1.json",
    "intraday_t": ROOT / "reports" / "research" / "intraday_t_v1_results.json",
    "info_prop": ROOT / "reports" / "research" / "information_propagation_no_if_primary_test.json",
}


def read_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as fh:
        return json.load(fh)


DATA = {name: read_json(path) for name, path in SOURCE_PATHS.items()}


def assert_close(actual: float, expected: float, tolerance: float = 1e-9) -> None:
    if not math.isclose(float(actual), float(expected), rel_tol=0.0, abs_tol=tolerance):
        raise AssertionError(f"数值断言失败：actual={actual}, expected={expected}")


goal_v2 = DATA["goal_v2"]
goal_v1 = DATA["goal_v1"]
micro = DATA["micro"]
micro_discovery = DATA["micro_discovery"]
registered = DATA["registered"]
phase1 = DATA["phase1"]
r6 = DATA["r6"]
overlay = DATA["overlay"]
t_only = DATA["t_only"]
cross_etf = DATA["cross_etf"]
atlas = DATA["atlas"]
promotion = DATA["promotion"]
intraday_t = DATA["intraday_t"]
info_prop = DATA["info_prop"]

# 核心一致性断言：生成失败比悄悄写错数字更安全。
assert goal_v2["goal_achieved"] is False
assert_close(goal_v2["success_contract"]["target_net_sharpe_zero_cash_rate"], 1.2)
assert micro["candidate_name"] == "MICRO_CONSENSUS3_OR_SHORT_EXTREMES_DUAL_SHADOW"
assert_close(micro["primary_base_start_summaries"][0]["net_sharpe_zero_cash_rate"], 1.5803113646)
assert_close(micro["primary_base_start_summaries"][0]["strategy_cagr"], 0.2654985684)
assert_close(overlay["results"]["baseline_base"]["sharpe_zero_cash_rate"], 1.42774159477833)
assert_close(overlay["results"]["overlay_base"]["sharpe_zero_cash_rate"], 1.257319126681689)
assert cross_etf["passing_candidates"] == []
assert atlas["adjudication"]["net_sharpe"] == "NOT_COMPUTED"
assert promotion["eligible_mechanism_count"] == 0
assert_close(intraday_t["summaries"]["BASE_COST"]["full_period"]["intraday_t"]["net_pnl_cny"], -5699.1, 1e-6)
assert_close(info_prop["base_round_trip_cost_bps"], 18.5)


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6573"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "2F5D50"
WHITE = "FFFFFF"
BLACK = "111111"
BODY_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


STATUS_CN = {
    "HISTORICAL_POST_SELECTION_AUDIT_REJECTED_NO_RESCUE": "历史合并期有效，但属于后选；不允许营救",
    "NOT_TARGET_QUALIFIED": "未达到目标资格",
    "NOT_ELIGIBLE": "不合格",
    "STRATEGY_GATE_FAILED": "策略门失败",
    "HISTORICAL_GATE_FAILED": "历史门失败",
    "RETROSPECTIVE_FAIL_AND_INSUFFICIENT_EVIDENCE_DO_NOT_INTEGRATE": "回溯增量失败，不整合",
    "REJECTED_FIXED_CROSS_ETF_FORCED_FLOW_BINARY_FAMILY_NO_RESCUE": "固定跨ETF资金流家族拒绝",
    "REJECTED_FROZEN_INSUFFICIENT_EVENT_SUPPORT_NO_RESCUE": "事件样本不足，冻结拒绝",
    "HISTORICAL_REJECTED_FROZEN": "历史拒绝并冻结",
    "FROZEN_CANDIDATE_REJECTED_CONTINUE_SEARCH": "候选拒绝，继续搜索",
    "FIXED_RULE_VALIDATION_REJECTED_NO_PARAMETER_RESCUE": "固定规则验证失败，不许调参营救",
    "DEVELOPMENT_DISCOVERY_COMPLETE_NO_2024_PLUS_READ": "开发发现完成，无2024年后验证结论",
    "DEVELOPMENT_DISCOVERY_COMPLETE_NO_VALIDATION_READ": "开发发现完成，无验证结论",
    "ELIGIBLE_FOR_COMBINATION_DISCUSSION": "仅获第一轮组合讨论资格",
    "DESCRIPTIVE_ONLY_NO_PARAMETER_RESCUE": "仅描述性，不调参营救",
    "DESCRIPTIVE_ONLY": "仅描述性",
    "STRUCTURAL": "结构性风险信息",
    "DATA_BLOCKED": "数据合同阻断",
}


def pct(value: Any, digits: int = 2) -> str:
    if value is None:
        return "未计算"
    return f"{float(value) * 100:.{digits}f}%"


def pp(value: Any, digits: int = 2, signed: bool = True) -> str:
    if value is None:
        return "未计算"
    prefix = "+" if signed and float(value) > 0 else ""
    return f"{prefix}{float(value) * 100:.{digits}f}个百分点"


def sharpe(value: Any, digits: int = 3) -> str:
    if value is None:
        return "未计算"
    return f"{float(value):.{digits}f}"


def money(value: Any, digits: int = 2) -> str:
    if value is None:
        return "未计算"
    return f"{float(value):,.{digits}f}元"


def integer(value: Any) -> str:
    if value is None:
        return "未计算"
    return f"{int(value):,}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_run_font(
    run,
    *,
    latin: str = BODY_FONT,
    east_asia: str = CJK_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, latin: str, east_asia: str, size: float, color: str, bold: bool = False) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (("w:ascii", latin), ("w:hAnsi", latin), ("w:eastAsia", east_asia), ("w:cs", latin)):
        rfonts.set(qn(key), value)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def ensure_style(doc: Document, name: str, style_type: WD_STYLE_TYPE = WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, latin=BODY_FONT, east_asia=CJK_FONT, size=11, color=BLACK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    pf.widow_control = True

    title = doc.styles["Title"]
    set_style_font(title, latin=BODY_FONT, east_asia=CJK_FONT, size=30, color=INK, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, latin=BODY_FONT, east_asia=CJK_FONT, size=15, color=DARK_BLUE)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(24)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.keep_with_next = True

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        set_style_font(style, latin=BODY_FONT, east_asia=CJK_FONT, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, latin=BODY_FONT, east_asia=CJK_FONT, size=11, color=BLACK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    table_text = ensure_style(doc, "Table Text")
    set_style_font(table_text, latin=BODY_FONT, east_asia=CJK_FONT, size=8.6, color=BLACK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.05
    table_text.paragraph_format.widow_control = True

    table_header = ensure_style(doc, "Table Header")
    set_style_font(table_header, latin=BODY_FONT, east_asia=CJK_FONT, size=8.6, color=INK, bold=True)
    table_header.paragraph_format.space_before = Pt(0)
    table_header.paragraph_format.space_after = Pt(2)
    table_header.paragraph_format.line_spacing = 1.05

    source_text = ensure_style(doc, "Source Text")
    set_style_font(source_text, latin=BODY_FONT, east_asia=CJK_FONT, size=8.5, color=MUTED)
    source_text.paragraph_format.space_before = Pt(4)
    source_text.paragraph_format.space_after = Pt(4)
    source_text.paragraph_format.line_spacing = 1.0

    lead = ensure_style(doc, "Lead")
    set_style_font(lead, latin=BODY_FONT, east_asia=CJK_FONT, size=12, color=INK, bold=True)
    lead.paragraph_format.space_before = Pt(0)
    lead.paragraph_format.space_after = Pt(8)
    lead.paragraph_format.line_spacing = 1.25

    small = ensure_style(doc, "Small Note")
    set_style_font(small, latin=BODY_FONT, east_asia=CJK_FONT, size=9, color=MUTED)
    small.paragraph_format.space_before = Pt(0)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.15


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9DEE7", size: int = 4) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def set_table_row_cant_split(row) -> None:
    """避免同一条数据记录被拆到两页，提升长表的可读性。"""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
    *,
    font_size: float = 8.6,
    center_columns: Iterable[int] = (),
    zebra: bool = False,
) -> Any:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"表格列宽之和必须为{TABLE_WIDTH_DXA}，实际为{sum(widths_dxa)}")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    center_set = set(center_columns)
    for col, text in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Header"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col in center_set else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, color=INK, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
    set_repeat_table_header(table.rows[0])
    set_table_row_cant_split(table.rows[0])

    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        set_table_row_cant_split(row)
        cells = row.cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col in center_set else WD_ALIGN_PARAGRAPH.LEFT
            parts = str(value).split("\n")
            for part_index, part in enumerate(parts):
                if part_index:
                    p.add_run().add_break()
                r = p.add_run(part)
                set_run_font(r, size=font_size, color=BLACK)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if zebra and row_index % 2 == 0:
                set_cell_shading(cell, "FAFBFC")
            set_cell_border(cell)

    apply_table_geometry(
        table,
        widths_dxa,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS_DXA,
    )
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def shade_paragraph(paragraph, fill: str, border_color: str = BLUE) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_color)
    borders.append(left)


def add_callout(doc: Document, label: str, text: str, *, kind: str = "info") -> None:
    color = {"info": BLUE, "success": GREEN, "warning": GOLD, "risk": RED}.get(kind, BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, CALLOUT, color)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=11, color=color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=BLACK)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> Any:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead) :])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc: Document, items: Sequence[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_numbers(doc: Document, items: Sequence[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)


def add_source(doc: Document, *paths: Path) -> None:
    rels = [str(path.relative_to(ROOT)).replace("\\", "/") for path in paths]
    p = doc.add_paragraph(style="Source Text")
    r = p.add_run("数据来源：" + "；".join(rels))
    set_run_font(r, size=8.5, color=MUTED)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.append(fld_char1)
    field_run._r.append(instr)
    field_run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hrun = hp.add_run("510300历史策略与回测总报告")
    set_run_font(hrun, size=9, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(84)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    kr = kicker.add_run("510300 STRATEGY RESEARCH")
    set_run_font(kr, size=10.5, color=GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("510300历史策略与\n全部回测结果报告")
    for run in title.runs:
        set_run_font(run, size=30, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    sr = subtitle.add_run("历史有效性、成本稳健性、失败分支与目标差距的完整说明")
    set_run_font(sr, size=15, color=DARK_BLUE)

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_before = Pt(10)
    scope.paragraph_format.space_after = Pt(50)
    rr = scope.add_run("评价口径：只使用截至2026年8月30日已经存在的历史数据；前向数据不作为有效性门槛")
    set_run_font(rr, size=10.5, color=MUTED, italic=True)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    dr = date.add_run("报告日期：2026年8月30日")
    set_run_font(dr, size=11, color=INK, bold=True)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("执行资产：510300.SH / 现金 | 研究性质：历史证据汇总")
    set_run_font(nr, size=9.5, color=MUTED)

    doc.add_page_break()


def add_navigation(doc: Document) -> None:
    doc.add_heading("阅读导航", level=1)
    add_callout(
        doc,
        "一句话结论",
        "按“只看过去数据”的口径，历史表现最有效的是微观结构风险开关策略；估值仓位基线在较短的两年窗口内次之。若要求开发完成后在独立历史后段继续复现，则目前没有一条策略完全合格。",
        kind="success",
    )
    items = [
        "一、执行摘要：先给结论和历史有效性分层",
        "二、评价口径：解释为什么高夏普不一定等于有效策略",
        "三、统一交易合同与可比性限制",
        "四、历史冠军：微观结构风险开关策略",
        "五、第二名：估值仓位基线及15分钟择时叠加",
        "六、全部十个登记因子回测结果",
        "七、四轮因子筛选、Phase 1、R6及其他策略家族",
        "八、日内做T与分钟级信息传播研究",
        "九、机制图谱H1-H4及统计结果",
        "十、最终历史排名、目标差距和准确表述",
        "附录：指标词典、状态词典和证据文件哈希",
    ]
    add_numbers(doc, items)


def section_executive_summary(doc: Document) -> None:
    doc.add_heading("一、执行摘要", level=1)
    p = doc.add_paragraph(style="Lead")
    r = p.add_run("历史冠军已经可以明确，但“历史赚钱”与“未经后选污染的稳定规律”必须分开表述。")
    set_run_font(r, size=12, color=INK, bold=True)

    rows = [
        ["A", "微观结构风险开关", "历史合并期有效", "净夏普1.580；双倍成本1.522-1.590；最大回撤-10.15%", "历史第一，但后选且2026边际衰减"],
        ["B", "估值仓位基线", "短窗口历史有效", "基础/压力夏普1.428/1.392；最大回撤-6.75%/-6.90%", "约2.03年，证据长度不足"],
        ["C", "H3下行压力因子", "风险预测有效", "未来10日系数-0.439个百分点，HAC p=0.018", "能识别风险，不是收益策略"],
        ["D", "价量、PE/PB、Phase 1、R6、Donchian、跨ETF资金流、日内T", "局部漂亮或整体失败", "局部夏普可高于1.2，但完整历史、成本、交易次数或稳定性失败", "不应称为有效策略"],
        ["E", "宏观事件、H4日内压力等", "未获准评价", "事件门或数据合同失败，收益/夏普未计算", "NOT_ALLOWED不等于零收益"],
    ]
    add_table(doc, ["等级", "对象", "历史判定", "关键结果", "准确结论"], rows, [650, 1750, 1500, 3000, 2460], font_size=8.2, center_columns=(0,), zebra=True)
    add_source(doc, SOURCE_PATHS["goal_v2"], SOURCE_PATHS["micro"], SOURCE_PATHS["overlay"], SOURCE_PATHS["atlas"])

    doc.add_heading("核心结论", level=2)
    add_bullets(
        doc,
        [
            "按历史合并收益、起点扰动、双倍成本和删除单年检查，微观结构风险开关是最强候选。",
            "它的主要价值来自2022-2024年的风险规避；2025年贡献减弱，2026年截至8月27日几乎没有新增超额。",
            "估值仓位基线在2024年7月至2026年8月期间达到夏普1.428，但样本仅约两年；加入15分钟择时后反而降低夏普、提高回撤。",
            "十个登记因子中没有最终冻结策略。价量确认只在伪样本外窗口表现较好，完整历史夏普仅0.217；PE/PB的高夏普由极少交易产生。",
            "Phase 1、R6、Donchian、跨ETF份额流和日内T均有明确失败证据；机制图谱只留下H3下行压力这一条结构性风险线索。",
        ],
    )


def section_methodology(doc: Document) -> None:
    doc.add_heading("二、评价口径：怎样判断“过去有效”", level=1)
    add_body(
        doc,
        "本报告遵循用户指定口径：不要求等待未来新增数据，不把252日Shadow或任何尚未发生的前向样本作为历史有效性的必要条件。但过去数据内部仍然分为开发、伪样本外、受污染回顾、完整历史和时间复制等不同证据等级。",
    )
    rows = [
        ["开发期", "用于提出或调整规则的历史区间", "只能说明规则怎样被找到，不能单独证明有效"],
        ["伪样本外", "在研究流程中被当作后段验证，但整体研究过程已接触历史", "比开发期更强，但不是严格未见数据"],
        ["受污染回顾", "规则形成后再次查看、讨论或用于其他研究的历史", "可描述表现，不再具备独立验证含义"],
        ["完整历史", "把全部可用历史拼接计算", "适合看长期净值和成本，不足以消除后选偏差"],
        ["历史时间复制", "用较晚的历史切片检查早期规律是否继续出现", "是本报告判断近期衰减的重要依据"],
        ["收益未获准", "数据门、事件门或点时合同失败", "必须写未计算；不能填0，也不能当作亏损"],
    ]
    add_table(doc, ["证据层", "普通话解释", "能证明什么"], rows, [1500, 3300, 4560], font_size=8.6, zebra=True)
    add_callout(
        doc,
        "判定原则",
        "本报告把“历史有效”定义为：在过去数据合并回测中，扣除成本后收益和风险指标明确改善，并能通过至少一组历史稳健性检查。它不等同于“未来仍会有效”。",
        kind="warning",
    )


def section_contract(doc: Document) -> None:
    doc.add_heading("三、统一目标与交易合同", level=1)
    c = goal_v2["success_contract"]
    rows = [
        ["执行资产", "510300.SH；其他指数、成分股、期货或期权只能作为观察信息，不能作为本目标的交易资产"],
        ["允许持仓", "510300.SH或人民币现金"],
        ["小账户合同", f"初始资金{money(c['initial_capital_cny'], 0)}；每手{integer(c['lot_size_shares'])}股"],
        ["基础成本", f"单边佣金率{pct(c['commission_rate_per_leg'])}，最低{money(c['minimum_commission_cny_per_leg'], 0)}；单边滑点{c['base_slippage_bps_per_leg']:.0f}BP"],
        ["夏普目标", f"零现金利率口径，净夏普率不低于{c['target_net_sharpe_zero_cash_rate']:.1f}"],
        ["超额目标", f"累计净超额不低于{pct(c['target_cumulative_net_excess'])}"],
        ["执行时点", "日度策略原则上T日收盘后生成状态，T+1交易日开盘执行；100股整数手、现金约束和分红再投资按各审计合同执行"],
    ]
    add_table(doc, ["项目", "统一说明"], rows, [2100, 7260], font_size=9.2)
    add_callout(
        doc,
        "可比性限制",
        "部分早期因子筛选使用约10万元权益基准，而微观结构和日内T使用2万元小账户。夏普、CAGR和回撤可比较，但绝对手续费金额不能跨项目直接横比。",
        kind="info",
    )
    add_source(doc, SOURCE_PATHS["goal_v2"], SOURCE_PATHS["goal_v1"])


def section_microstructure(doc: Document) -> None:
    doc.add_heading("四、历史冠军：微观结构风险开关策略", level=1)
    add_callout(
        doc,
        "历史判定",
        "这是过去数据中表现最强的策略：完整合并期、五个起点和双倍成本下均超过夏普1.2。但它是历史后选序列，且2026年时间切片没有继续创造超额。准确称呼应为“历史有效、后选、近期衰减”。",
        kind="success",
    )

    doc.add_heading("4.1 普通人可以怎样理解这套策略", level=2)
    add_body(
        doc,
        "这不是预测每天涨跌的技术指标策略。它默认持有510300；当ETF折溢价、份额申赎、融资融券和融券库存同时出现异常时，判断市场可能处于资金压力状态，从而暂时转为现金。风险信号解除后，再回到510300。",
    )
    add_bullets(
        doc,
        [
            "七个微观领域中至少三个同时触发风险票，或者融券库存进入极端枯竭/拥挤，形成原始空仓条件。",
            "所有分位阈值只使用滚动历史：排名窗口252个信号日，最少需要126日。",
            "风险条件使用迟滞退出，不在阈值附近频繁来回切换。",
            "双影子门检查相同风险信号在最近242日是否仍有正贡献、最近60日是否不低于-1%；影子门关闭时不执行空仓。",
            "信号在T日收盘后确定，T+1开盘执行；只做多或现金，没有杠杆、做空和衍生品执行。",
        ],
    )

    doc.add_heading("4.2 详细触发条件", level=2)
    rule = micro_discovery["rule"]
    rule_rows = [[name, description] for name, description in rule["micro_domains"].items()]
    rule_rows.extend(
        [
            ["融券库存枯竭", "20日融券余量/融资融券余额的252日分位不高于1%，当日触发"],
            ["融券库存拥挤", "同一比例的252日分位达到95%进入，回落到70%退出"],
            ["最终现金状态", "微观风险票至少3票，或融券库存极端任一触发，并且242日/60日双影子门仍开启"],
        ]
    )
    add_table(doc, ["信号域", "冻结规则"], rule_rows, [2800, 6560], font_size=8.8, zebra=True)

    doc.add_heading("4.3 主历史结果", level=2)
    m = micro["primary_base_start_summaries"][0]
    d = micro["primary_double_cost_start_summaries"][0]
    main_rows = [
        ["评价区间", f"{m['start_date']}至{m['end_date']}", f"{d['start_date']}至{d['end_date']}"],
        ["净夏普率", sharpe(m["net_sharpe_zero_cash_rate"]), sharpe(d["net_sharpe_zero_cash_rate"])],
        ["CAGR", pct(m["strategy_cagr"]), pct(d["strategy_cagr"])],
        ["相对H00300年化超额", pp(m["annualized_excess_vs_h00300"]), pp(d["annualized_excess_vs_h00300"])],
        ["相对510300买入持有年化超额", pp(m["annualized_excess_vs_buy_hold"]), pp(d["annualized_excess_vs_buy_hold"])],
        ["最大回撤", pct(m["maximum_drawdown"]), pct(d["maximum_drawdown"])],
        ["买入持有最大回撤", pct(m["buy_hold_maximum_drawdown"]), pct(d["buy_hold_maximum_drawdown"])],
        ["现金日占比", pct(m["cash_day_share"]), pct(d["cash_day_share"])],
        ["交易腿数", integer(m["trade_leg_count"]), integer(d["trade_leg_count"])],
        ["单年最多完整往返", integer(m["maximum_complete_round_trips_in_any_year"]), integer(d["maximum_complete_round_trips_in_any_year"])],
        ["佣金", money(m["total_commission_cny"]), money(d["total_commission_cny"])],
        ["滑点成本", money(m["total_slippage_cost_cny"]), money(d["total_slippage_cost_cny"])],
    ]
    add_table(doc, ["指标", "基础成本", "双倍成本"], main_rows, [3300, 3030, 3030], font_size=9.0, center_columns=(1, 2), zebra=True)

    heading = doc.add_heading("4.4 五个起点与成本稳健性", level=2)
    heading.paragraph_format.page_break_before = True
    start_rows = []
    for base, double in zip(micro["primary_base_start_summaries"], micro["primary_double_cost_start_summaries"]):
        start_rows.append(
            [
                f"+{base['start_perturbation']}日",
                base["first_execution_date"],
                sharpe(base["net_sharpe_zero_cash_rate"]),
                pct(base["strategy_cagr"]),
                pp(base["annualized_excess_vs_h00300"]),
                pct(base["maximum_drawdown"]),
                sharpe(double["net_sharpe_zero_cash_rate"]),
            ]
        )
    add_table(
        doc,
        ["起点", "首次执行", "基础夏普", "基础CAGR", "对H00300超额", "基础回撤", "双倍成本夏普"],
        start_rows,
        [720, 1350, 1120, 1050, 1800, 1250, 2070],
        font_size=8.2,
        center_columns=(0, 1, 2, 3, 5, 6),
        zebra=True,
    )

    doc.add_heading("4.5 删除单年与年份贡献", level=2)
    loo = micro["year_robustness_primary_start0"]["leave_one_calendar_year_out_sharpes"]
    contrib = micro["year_robustness_primary_start0"]["yearly_log_excess_contributions_vs_buy_hold"]
    year_rows = [[year, sharpe(loo[year]), pct(contrib[year]), "对数超额贡献，不是当年简单收益率"] for year in loo]
    add_table(doc, ["被删除年份", "删除后夏普", "该年对数超额贡献", "解释"], year_rows, [1300, 1500, 2100, 4460], font_size=8.7, center_columns=(0, 1, 2), zebra=True)
    add_body(
        doc,
        f"删除任意一个完整年份后，夏普最低仍为{min(float(v) for v in loo.values()):.3f}。单一年份占总对数超额的最大比例为{pct(micro['year_robustness_primary_start0']['maximum_single_year_share_of_total_log_excess'])}，未超过50%。这说明合并历史并非完全由单一年份制造。",
    )

    doc.add_heading("4.6 2026年历史时间切片：近期边际已经衰减", level=2)
    temporal_rows = []
    for item in micro["temporal_replication_base_start_summaries"]:
        temporal_rows.append(
            [
                f"+{item['start_perturbation']}日",
                sharpe(item["net_sharpe_zero_cash_rate"]),
                pct(item["strategy_cagr"]),
                pp(item["annualized_excess_vs_h00300"]),
                pp(item["annualized_excess_vs_buy_hold"]),
                pct(item["cash_day_share"]),
            ]
        )
    add_table(doc, ["起点", "夏普", "CAGR", "对H00300超额", "对买入持有超额", "现金日占比"], temporal_rows, [900, 1100, 1200, 2050, 2050, 2060], font_size=8.5, center_columns=(0, 1, 2, 5), zebra=True)
    add_callout(
        doc,
        "准确解释",
        "2026年切片已经属于过去数据。策略基本一直持有510300，没有创造买入持有之外的价值；因此应当说它“历史合并期有效、近期失去增量”，不能说“每个年份都有效”。",
        kind="warning",
    )
    add_source(doc, SOURCE_PATHS["micro"], SOURCE_PATHS["micro_discovery"])


def section_valuation(doc: Document) -> None:
    doc.add_heading("五、第二名：估值仓位基线与15分钟择时叠加", level=1)
    add_body(
        doc,
        "估值仓位基线根据估值模型调整目标仓位；15分钟叠加层只决定普通买卖何时执行，不改变估值目标仓位，危机减仓不等待技术信号。历史结果表明，估值仓位本身在两年窗口内有效，但15分钟择时没有增加价值。",
    )
    result_rows = []
    labels = {
        "baseline_base": "估值基线-基础成本",
        "baseline_stress": "估值基线-压力成本",
        "overlay_base": "15分钟叠加-基础成本",
        "overlay_stress": "15分钟叠加-压力成本",
    }
    for key in ("baseline_base", "baseline_stress", "overlay_base", "overlay_stress"):
        item = overlay["results"][key]
        result_rows.append(
            [
                labels[key],
                f"{item['start_date']}至{item['end_date']}",
                pct(item["total_return"]),
                pct(item["cagr"]),
                sharpe(item["sharpe_zero_cash_rate"]),
                pct(item["max_drawdown"]),
                pct(item["average_position"]),
                integer(item["trade_count"]),
            ]
        )
    add_table(doc, ["版本", "区间", "总收益", "CAGR", "夏普", "最大回撤", "平均仓位", "交易"], result_rows, [1850, 1800, 900, 850, 720, 1050, 1050, 1140], font_size=8.0, center_columns=(2, 3, 4, 5, 6, 7), zebra=True)
    add_bullets(
        doc,
        [
            f"基础成本下，叠加层使夏普从{sharpe(overlay['results']['baseline_base']['sharpe_zero_cash_rate'])}降到{sharpe(overlay['results']['overlay_base']['sharpe_zero_cash_rate'])}。",
            f"最大回撤从{pct(overlay['results']['baseline_base']['max_drawdown'])}恶化到{pct(overlay['results']['overlay_base']['max_drawdown'])}。",
            f"叠加层结束资金比基线少{money(abs(overlay['relative_diagnostics']['ending_equity_difference_cny']))}。",
            "前后半段相对收益一正一负，9项增量检查全部未通过。",
        ],
    )
    add_callout(
        doc,
        "历史判定",
        "估值基线可列为“短窗口历史有效”；15分钟择时叠加明确无效。不能用叠加后的1.229压力夏普来宣称技术择时有效，因为它低于不使用技术择时的基线。",
        kind="warning",
    )
    add_source(doc, SOURCE_PATHS["overlay"])


def factor_metric_cell(block: dict[str, Any]) -> str:
    s = block["strategy"]
    return (
        f"夏普 {sharpe(s.get('sharpe_zero_cash_rate'))}\n"
        f"CAGR {pct(s.get('cagr'))}\n"
        f"超额 {pp(block.get('excess_vs_realistic_buy_hold'))}\n"
        f"回撤 {pct(s.get('max_drawdown'))}\n"
        f"交易 {integer(s.get('trade_count'))}"
    )


def section_registered_factors(doc: Document) -> None:
    doc.add_heading("六、全部十个登记因子回测结果", level=1)
    add_body(
        doc,
        "第一轮对十个登记因子采用统一长仓/现金规则：进入分位70%、退出分位30%、中间维持原状态，最少252日历史、756日分位窗口。基础成本为5BP滑点，压力成本为15BP滑点。治理标签是“回溯+伪样本外，无严格未见留出”。",
    )

    sharpe_rows = []
    detail_rows = []
    for factor in registered["factors"]:
        periods = factor["periods"]
        stress = factor["pseudo_oos_stress_15bps"]
        status = STATUS_CN.get(factor["candidate_status"], factor["candidate_status"])
        suff = factor["posthoc_sample_sufficiency_guard"]
        sharpe_rows.append(
            [
                factor["factor_name_cn"],
                sharpe(periods["development"]["strategy"]["sharpe_zero_cash_rate"]),
                sharpe(periods["pseudo_oos"]["strategy"]["sharpe_zero_cash_rate"]),
                sharpe(stress["strategy"]["sharpe_zero_cash_rate"]),
                sharpe(periods["contaminated_retrospective"]["strategy"]["sharpe_zero_cash_rate"]),
                sharpe(periods["full_retrospective"]["strategy"]["sharpe_zero_cash_rate"]),
                f"{status}\n样本门：{'通过' if suff['passes'] else '失败'}（{suff['observed_trade_count']}笔）",
            ]
        )
        detail_rows.append(
            [
                factor["factor_name_cn"],
                factor_metric_cell(periods["development"]),
                factor_metric_cell(periods["pseudo_oos"]),
                factor_metric_cell(stress),
                factor_metric_cell(periods["full_retrospective"]),
            ]
        )

    add_table(
        doc,
        ["因子", "开发夏普", "伪样本外", "15BP压力", "受污染后段", "完整历史", "第一轮状态/样本门"],
        sharpe_rows,
        [2500, 900, 1000, 1000, 1100, 1000, 1860],
        font_size=7.9,
        center_columns=(1, 2, 3, 4, 5),
        zebra=True,
    )
    add_source(doc, SOURCE_PATHS["registered"])

    heading = doc.add_heading("6.1 十个因子的详细收益、超额、回撤与交易次数", level=2)
    heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["因子", "开发期", "伪样本外", "15BP压力", "完整历史"],
        detail_rows,
        [1900, 1865, 1865, 1865, 1865],
        font_size=7.6,
        center_columns=(1, 2, 3, 4),
        zebra=True,
    )

    doc.add_heading("6.2 逐项结论", level=2)
    add_bullets(
        doc,
        [
            "510300过去5日收益：完整历史夏普0.229；伪样本外跑输买入持有16.20个百分点，不合格。",
            "沪深300过去20日收益：完整历史夏普0.268；伪样本外和压力期均跑输买入持有，不合格。",
            "20/60日均线、60日区间位置和20日实现波动率：伪样本外夏普为负或接近零，完整历史也没有稳定优势。",
            "短中期波动率比：伪样本外夏普0.686，但伪样本外、压力期和后段均跑输买入持有；完整历史夏普0.310。",
            "价格与成交量确认：伪样本外夏普1.376，是第一轮唯一临时候选；15BP压力降至1.146，后段跑输买入持有2.00个百分点，完整历史夏普0.217，因此被第二轮退役。",
            "收盘价相对VWAP：完整历史夏普-0.475，且高换手、成本后明显失败。",
            "PE五年分位：伪样本外夏普2.216，但只有4笔交易；后段仅2笔且跑输买入持有7.27个百分点，不通过最低交易次数。",
            "PB五年分位：伪样本外只有2笔交易；后段夏普1.613也只有2笔且跑输买入持有3.98个百分点。",
        ],
    )

    heading = doc.add_heading("6.3 四轮因子筛选总结果", level=2)
    heading.paragraph_format.page_break_before = True
    round_rows = [
        ["Round 1", "10个日度因子+6个日内假设", "价量确认仅为临时讨论候选；PE/PB交易次数不足；无策略冻结"],
        ["Round 2", "价量机制拆解、量冲击、IF相对领先", "没有因子通过伪样本外超额和15BP压力门；价量临时候选退役"],
        ["Round 3", "6个等权宽度/市值参与因子", "全部未通过D20、D5、伪样本外、风险和成本的联合门槛"],
        ["Round 4", "6个流通市值近似加权宽度因子", "全部伪样本外及15BP压力超额为负；近似权重不得冒充官方权重"],
    ]
    add_table(doc, ["轮次", "研究范围", "最终结果"], round_rows, [1300, 3000, 5060], font_size=8.8, zebra=True)
    add_source(doc, SOURCE_PATHS["round1"], SOURCE_PATHS["round2"], SOURCE_PATHS["round3"], SOURCE_PATHS["round4"])


def section_other_strategies(doc: Document) -> None:
    doc.add_heading("七、Phase 1、R6与其他策略家族", level=1)

    doc.add_heading("7.1 Phase 1估值迟滞模型", level=2)
    phase_rows = []
    phase_labels = {
        "development": "开发期",
        "validation": "顺序验证期",
        "retrospective_test": "受污染回顾期",
        "full_five_years": "完整五年",
    }
    for key in ("development", "validation", "retrospective_test", "full_five_years"):
        block = phase1["periods"][key]
        s = block["strategy"]
        phase_rows.append(
            [
                phase_labels[key],
                f"{s['start_date']}至{s['end_date']}",
                pct(s["total_return"]),
                sharpe(s["sharpe_zero_cash_rate"]),
                pct(s["max_drawdown"]),
                integer(s["trade_count"]),
                pp(block.get("excess_vs_etf_total_return")),
            ]
        )
    add_table(doc, ["区间", "日期", "总收益", "夏普", "最大回撤", "交易", "对ETF总超额"], phase_rows, [1200, 1850, 1000, 900, 1200, 800, 2410], font_size=8.5, center_columns=(2, 3, 4, 5), zebra=True)
    add_body(
        doc,
        "顺序验证期夏普2.373看起来非常高，但只有4笔交易；开发期夏普-0.323，受污染回顾期0笔交易，完整五年夏普0.356。自然年度没有一年达到20个百分点超额，策略门失败。",
    )

    doc.add_heading("7.2 R6状态配置家族", level=2)
    r6_find = next(item for item in goal_v2["registered_machine_result_inventory"]["findings"] if item["candidate"] == "R6_MODEL_FAMILY")
    add_bullets(
        doc,
        [
            f"共测试{r6['candidate_count']}个冻结候选，历史门通过数为{r6['historical_gate_pass_count']}。",
            f"最高的单个2025年夏普为{sharpe(r6_find['maximum_single_2025_sharpe'])}，但同期跑输ETF {abs(float(r6_find['associated_2025_excess_vs_etf'])) * 100:.2f}个百分点。",
            f"对应历史伪样本外夏普为{sharpe(r6_find['associated_historical_pseudo_oos_sharpe'])}，并跑输ETF {abs(float(r6_find['associated_historical_pseudo_oos_excess_vs_etf'])) * 100:.2f}个百分点。",
            "所有邻域组正超额比例为0；最终属于负向动态择时案例，不是高夏普策略。",
        ],
    )

    doc.add_heading("7.3 T-only Donchian历史稳健性", level=2)
    base_t = t_only["scenarios"]["BASELINE_REPLICATION"]["metrics"]
    t_rows = [
        ["基线", sharpe(base_t["sharpe"]), pct(base_t["cagr"]), pct(base_t["maximum_drawdown"]), pct(base_t["average_exposure"]), money(base_t["ending_equity_cny"])],
        ["延迟1根K线", sharpe(t_only["scenarios"]["DELAY_1_EXTRA_BAR"]["metrics"]["sharpe"]), pct(t_only["scenarios"]["DELAY_1_EXTRA_BAR"]["metrics"]["cagr"]), pct(t_only["scenarios"]["DELAY_1_EXTRA_BAR"]["metrics"]["maximum_drawdown"]), pct(t_only["scenarios"]["DELAY_1_EXTRA_BAR"]["metrics"]["average_exposure"]), money(t_only["scenarios"]["DELAY_1_EXTRA_BAR"]["metrics"]["ending_equity_cny"])],
        ["双倍总成本", sharpe(t_only["double_total_cost"]["sharpe"]), pct(t_only["double_total_cost"]["cagr"]), pct(t_only["double_total_cost"]["maximum_drawdown"]), pct(t_only["double_total_cost"]["average_exposure"]), money(t_only["double_total_cost"]["ending_equity_cny"])],
    ]
    add_table(doc, ["情景", "夏普", "CAGR", "最大回撤", "平均敞口", "期末权益"], t_rows, [1900, 1200, 1200, 1500, 1600, 1960], font_size=8.8, center_columns=(1, 2, 3, 4, 5), zebra=True)
    add_body(doc, "全样本净夏普0.405，明显低于1.2；增量Bootstrap区间跨过零，且历史从2012年开始，不符合本轮2021年后目标合同。结论为仅描述性，不调参营救。")

    doc.add_heading("7.4 跨ETF官方份额强制资金流", level=2)
    best = cross_etf["best_candidate"]
    add_bullets(
        doc,
        [
            f"固定5个候选、2套成本、3个时期、5个起点，共{cross_etf['result_rows']}条结果。",
            f"通过候选数为{len(cross_etf['passing_candidates'])}；可进入Shadow的候选数为{len(cross_etf['positive_shadow_candidates'])}。",
            f"最好候选{best['candidate_id']}的最差压力年化超额为{pp(best['stress_min_annualized_excess'])}，后期方向反转。",
            "正式结论：整个固定家族拒绝，不改阈值、周数、ETF池、执行延迟或组合条件营救。",
        ],
    )

    doc.add_heading("7.5 其他已登记候选与未运行分支", level=2)
    excluded = goal_v1["candidate_screen"]["excluded_candidates"]
    other_rows = [
        [item["candidate"], STATUS_CN.get(item["status"], item["status"]), item["reason"]]
        for item in excluded
    ]
    macro = goal_v1["historical_adjudication"]["macro_stress_avoidance"]
    other_rows.insert(
        0,
        ["MACRO_STRESS_AVOIDANCE", STATUS_CN.get(macro["status"], macro["status"]), macro["interpretation"]],
    )
    add_table(doc, ["候选", "状态", "普通话解释"], other_rows, [2450, 2600, 4310], font_size=7.9, zebra=True)
    add_callout(
        doc,
        "状态语义",
        "宏观压力分支完整事件为0、最低要求8，因此组合回测被跳过。NOT_ALLOWED既不是零收益，也不是亏损；它表示数据或事件门没有授权计算。",
        kind="info",
    )
    add_source(doc, SOURCE_PATHS["phase1"], SOURCE_PATHS["r6"], SOURCE_PATHS["t_only"], SOURCE_PATHS["cross_etf"], SOURCE_PATHS["goal_v1"])


def section_intraday(doc: Document) -> None:
    doc.add_heading("八、日内做T与分钟级信息传播", level=1)

    doc.add_heading("8.1 VWAP库存均值回归做T", level=2)
    base = intraday_t["summaries"]["BASE_COST"]["full_period"]
    t = base["intraday_t"]
    intraday_rows = [
        ["完整区间", f"{base['actual_start'][:10]}至{base['actual_end'][:10]}"],
        ["策略总收益", pct(base["strategy"]["total_return"])],
        ["策略夏普", sharpe(base["strategy"]["sharpe_zero_cash_rate"])],
        ["完整往返次数", integer(t["round_trip_count"])],
        ["胜率", pct(t["win_rate"])],
        ["原始毛利润", money(t["gross_raw_pnl_cny"])],
        ["执行后毛利润", money(t["execution_gross_pnl_cny"])],
        ["佣金", money(t["commission_cny"])],
        ["滑点与最小价位成本", money(t["slippage_and_tick_cost_cny"])],
        ["做T净利润", money(t["net_pnl_cny"])],
    ]
    add_table(doc, ["指标", "结果"], intraday_rows, [3600, 5760], font_size=9.0, center_columns=(1,), zebra=True)
    add_body(doc, "原始价差已经略为负，扣除执行摩擦后净亏损5,699.10元。问题不是参数不够好，而是5年内312次小额往返在最低佣金和滑点下缺乏经济空间。")

    doc.add_heading("8.2 分钟级信息传播：有统计排序，但没有可交易Alpha", level=2)
    dev = info_prop["periods"]["development"]
    val = info_prop["periods"]["validation"]
    info_rows = [
        ["开发期-收盘后5分钟标记", integer(dev["etf_markout_close_5m"]["observation_count"]), f"{dev['etf_markout_close_5m']['p10_minus_p1_bps']:.2f}BP", "统计排序存在"],
        ["开发期-下一开盘后5分钟", integer(dev["etf_execution_next_open_5m"]["observation_count"]), f"{dev['etf_execution_next_open_5m']['p10_minus_p1_bps']:.2f}BP", "低于18.5BP成本"],
        ["验证期-收盘后5分钟标记", integer(val["etf_markout_close_5m"]["observation_count"]), f"{val['etf_markout_close_5m']['p10_minus_p1_bps']:.2f}BP", "统计排序复现"],
        ["验证期-下一开盘后5分钟", integer(val["etf_execution_next_open_5m"]["observation_count"]), f"{val['etf_execution_next_open_5m']['p10_minus_p1_bps']:.2f}BP", f"P10成本后均值{val['etf_execution_next_open_5m_net_base']['p10_mean_bps']:.2f}BP"],
    ]
    add_table(doc, ["测试", "样本", "P10-P1", "解释"], info_rows, [3000, 1200, 1500, 3660], font_size=8.7, center_columns=(1, 2), zebra=True)
    add_callout(
        doc,
        "结论",
        "分钟级价格发现顺序是真实的统计关系，但验证期可执行毛差只有1.79BP，远低于18.5BP基础往返成本；因此被拒绝为独立交易Alpha。",
        kind="risk",
    )
    add_source(doc, SOURCE_PATHS["intraday_t"], SOURCE_PATHS["info_prop"])


def section_atlas(doc: Document) -> None:
    doc.add_heading("九、机制图谱H1-H4：哪些规律值得保留", level=1)
    panel = atlas["state_panel"]
    add_body(
        doc,
        f"机制图谱覆盖{panel['first_date']}至{panel['last_date']}，共{integer(panel['rows'])}个交易日、{integer(panel['columns'])}个状态字段；建立6个机制病例和18个不复用失败对照，并运行5/10/20/60日局部投影。该项目只允许机制发现，不运行组合回测。",
    )

    doc.add_heading("9.1 四个机制的最终状态", level=2)
    mech_rows = [
        ["H1 趋势扩散", "慢趋势、内部宽度扩散、风险未加速、IF未恶化", "MACD和宽度方向不稳定，统计证据不足", "仅描述性"],
        ["H2 真假修复", "MACD负值收敛时，宽度、风险降温和IF是否确认", "确认状态仅9/1211日；10日均值1.90%，比未确认高1.137个百分点", "稀疏、描述性"],
        ["H3 去杠杆压力", "下行压力上升是否预示未来收益和路径恶化", "10日系数-0.439个百分点，p=0.018；前后时期方向一致", "结构性风险信息"],
        ["H4 日内暂时压力", "ETF-IF-IOPV偏离是否日内收敛", "缺少覆盖六病例和长期对照的同步历史", "数据合同失败"],
    ]
    add_table(doc, ["机制", "研究问题", "历史结果", "当前身份"], mech_rows, [1600, 2900, 3100, 1760], font_size=8.4, zebra=True)

    doc.add_heading("9.2 连续条件响应结果", level=2)
    factor_rows = []
    for item in atlas["factor_matrix"]:
        factor_rows.append(
            [
                item["factor"],
                item["economic_mechanism"],
                str(item["primary_horizon_trading_days"]),
                "未估计" if item["full_beta"] is None else f"{float(item['full_beta']) * 100:+.3f}个百分点",
                "未估计" if item["full_p_value_hac"] is None else f"{float(item['full_p_value_hac']):.3f}",
                "是" if item["direction_consistent_early_late"] else "否",
                STATUS_CN.get(item["current_status"], item["current_status"]),
            ]
        )
    add_table(doc, ["因子", "机制", "主期限", "标准化系数", "HAC p值", "前后同向", "结论"], factor_rows, [1800, 2000, 800, 1450, 1000, 1000, 1310], font_size=7.7, center_columns=(2, 3, 4, 5), zebra=True)

    doc.add_heading("9.3 机会频率与条件收益", level=2)
    opp_rows = []
    state_cn = {
        "M_NEGATIVE_HIST_CONVERGING": "MACD负柱收敛",
        "BREADTH_IMPROVING": "宽度改善",
        "DOWNSIDE_COOLING": "下行风险降温",
        "IF_BASIS_IMPROVING": "IF基差改善",
        "H2_CONFIRMED": "H2全部确认",
        "H2_UNCONFIRMED": "H2未全部确认",
    }
    for item in atlas["opportunity_frequency_all"]:
        opp_rows.append(
            [
                state_cn.get(item["state"], item["state"]),
                integer(item["opportunity_days"]),
                pct(item["opportunity_frequency"]),
                pct(item["forward_return_10d_mean"]),
                pct(item["forward_return_20d_mean"]),
                "描述性，不是策略收益",
            ]
        )
    add_table(doc, ["状态", "天数", "频率", "后10日均值", "后20日均值", "解释"], opp_rows, [1900, 900, 1100, 1400, 1400, 2660], font_size=8.5, center_columns=(1, 2, 3, 4), zebra=True)

    doc.add_heading("9.4 病例与失败对照的差异", level=2)
    cc = atlas["case_control_comparison"]
    cc_rows = []
    for horizon in ("horizon_5d", "horizon_10d", "horizon_20d"):
        item = cc[horizon]
        label = horizon.replace("horizon_", "").upper()
        cc_rows.append([label, integer(item["case_count"]), integer(item["control_count"]), pct(item["case_mean"]), pct(item["control_mean"]), pp(item["case_minus_control_mean"])])
    add_table(doc, ["期限", "病例", "对照", "病例均值", "对照均值", "病例-对照"], cc_rows, [1200, 1000, 1000, 1700, 1700, 2760], font_size=8.8, center_columns=(0, 1, 2, 3, 4, 5), zebra=True)
    add_callout(
        doc,
        "重要偏差",
        "失败对照按未来10日收益不大于0筛选，因此病例-对照差异只能用于机制尸检，不能解释为无偏效应、因果效应或可交易收益。",
        kind="warning",
    )
    add_callout(
        doc,
        "机制结论",
        "H3下行压力是唯一值得保留的结构性风险因子，但它尚未形成买入、减仓、持有和退出规则；机制图谱的净夏普和累计净超额均未计算。",
        kind="info",
    )
    add_source(doc, SOURCE_PATHS["atlas"], SOURCE_PATHS["promotion"])


def section_final(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("十、最终历史排名、目标差距与准确表述", level=1)

    doc.add_heading("10.1 历史有效性排名", level=2)
    rows = [
        ["1", "微观结构风险开关", "历史合并期有效", "夏普1.580；双倍成本仍>1.5；回撤显著下降", "后选，2026无新增超额"],
        ["2", "估值仓位基线", "短窗口历史有效", "两年基础/压力夏普1.428/1.392", "历史长度短，无独立后段证明"],
        ["3", "H3下行压力", "风险信息有效", "5/10日负向系数显著或接近显著，方向稳定", "不是收益策略"],
        ["4", "价量确认20日", "阶段性有效", "伪样本外夏普1.376", "压力<1.2，后段跑输，完整历史0.217"],
        ["5", "PE/PB、Phase 1", "高夏普但样本稀疏", "局部夏普可达1.6-2.5", "2-4笔交易或后段0笔，不可信"],
        ["6", "R6、Donchian、跨ETF份额流、日内T", "历史无效", "风险可能下降，但净超额、成本或完整历史失败", "冻结拒绝"],
        ["7", "宏观事件、H4", "不可评价", "收益/夏普未计算", "数据或事件门失败"],
    ]
    add_table(doc, ["排名", "对象", "历史身份", "为什么排在这里", "关键限制"], rows, [700, 2000, 1500, 2900, 2260], font_size=8.2, center_columns=(0,), zebra=True)

    doc.add_heading("10.2 与夏普1.2目标的差距", level=2)
    gap_rows = [
        ["纯历史合并口径", "净夏普>=1.2", "微观结构1.580", "数值上超过0.380"],
        ["双倍成本口径", "仍能承受明显更高成本", "微观结构1.522-1.590", "仍超过目标"],
        ["短窗口历史口径", "净夏普>=1.2", "估值基线压力1.392", "数值达到，但仅约2年"],
        ["干净历史时间复制", "开发后独立后段仍复现", "微观结构2026夏普-0.099至0.134；其他候选也失败", "没有合格策略"],
        ["机制闭环", "经济机制+规则+成本后组合收益", "只有H3风险因子，尚无收益规则", "仍缺完整策略"],
    ]
    add_table(doc, ["评价层", "目标", "当前", "差距"], gap_rows, [1900, 2500, 3000, 1960], font_size=8.7, zebra=True)

    doc.add_heading("10.3 应当怎样对外表述", level=2)
    add_bullets(
        doc,
        [
            "可以说：‘微观结构风险开关在2022年3月至2026年8月的历史合并回测中，基础成本净夏普1.580，双倍成本仍超过1.5。’",
            "必须同时说：‘该序列属于历史后选，2026年时间切片没有继续创造超额，因此不能证明当前或未来仍有效。’",
            "可以说：‘估值仓位基线在约两年的历史窗口中达到夏普1.428；加入15分钟择时后表现变差。’",
            "不能说：‘我们已经拥有一条经过独立历史验证、稳定夏普1.2以上的策略。’",
            "不能把NOT_ALLOWED写成0收益，也不能把2-4笔交易产生的高夏普写成稳定Alpha。",
        ],
    )
    add_callout(
        doc,
        "最终结论",
        "如果只问过去哪条策略最有效，答案是微观结构风险开关；如果问过去数据能否证明某条策略具有不依赖后选的稳定预测力，答案仍是没有。",
        kind="success",
    )


def section_appendices(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("附录A：常用指标词典", level=1)
    glossary = [
        ["净夏普率", "策略年化平均收益除以年化波动率，本报告现金利率按0；越高代表单位波动获得的收益越多"],
        ["CAGR", "复合年化增长率；回答长期平均每年增长多少，不代表每年都取得相同收益"],
        ["最大回撤", "从历史净值高点到随后低点的最大跌幅；越接近0越好"],
        ["年化超额", "策略CAGR减去基准CAGR；与累计超额不是同一个量"],
        ["交易腿", "一次买入或卖出算一条腿；完整往返一般包含买入和卖出"],
        ["起点扰动", "把开始执行日向后移动若干交易日，检查结果是否依赖某个幸运起点"],
        ["双倍成本", "把佣金、最低佣金和滑点提高，用来检查策略是否只靠低估交易摩擦才赚钱"],
        ["伪样本外", "流程上像验证期，但研究整体已接触历史；强于开发期，弱于真正未见数据"],
        ["后选", "先看过历史结果，再形成或挑选规则；历史表现容易被高估"],
        ["HAC p值", "考虑时间序列相关和异方差后的显著性指标；越小表示随机波动解释该结果的可能性越低"],
    ]
    add_table(doc, ["术语", "普通话解释"], glossary, [1900, 7460], font_size=9.0, zebra=True)

    doc.add_heading("附录B：状态词典", level=1)
    states = [
        ["HISTORICAL_EFFECTIVE", "过去合并回测中成本后收益、风险和稳健性达到既定标准；不自动代表未来有效"],
        ["DESCRIPTIVE_ONLY", "能描述历史关系，但没有获准成为收益策略"],
        ["STRUCTURAL_RISK_FACTOR", "方向具有一定跨期一致性，可作为风险信息；仍缺交易规则"],
        ["NOT_ELIGIBLE / REJECTED", "候选未通过冻结门槛；不得用同一历史改参数营救"],
        ["NOT_ALLOWED", "数据、事件或治理门失败，收益不允许计算；不是0或亏损"],
        ["SKIPPED", "按预先规则跳过计算；没有产生组合回测结果"],
        ["NO_VIEW", "证据合同不完整，不能对结果作方向判断"],
    ]
    add_table(doc, ["状态", "含义"], states, [3000, 6360], font_size=8.9, zebra=True)

    doc.add_heading("附录C：证据文件与SHA-256", level=1)
    source_rows = []
    for name, path in SOURCE_PATHS.items():
        source_rows.append(
            [
                name,
                str(path.relative_to(ROOT)).replace("\\", "/"),
                f"{path.stat().st_size:,}",
                sha256(path),
            ]
        )
    add_table(doc, ["标识", "相对路径", "字节", "SHA-256"], source_rows, [1100, 3700, 900, 3660], font_size=6.8, center_columns=(2,), zebra=True)
    add_body(
        doc,
        "这些哈希用于证明本报告引用的是生成时读取的具体文件版本。若任何源文件后续变化，应重新生成本报告，而不是手工修改表格数字。",
    )

    doc.add_heading("附录D：报告范围与排除项", level=1)
    add_bullets(
        doc,
        [
            "纳入：与510300夏普1.2目标直接相关的历史策略、登记因子、成本压力、时间切片、机制研究、失败分支和日内研究。",
            "不纳入有效性判定：尚未发生的未来数据、未来Shadow日数和任何计划任务状态。",
            "不把其他股票、数字资产或A股横截面策略的收益借给510300。",
            "不把指数观察变量、IF信息或跨ETF份额代理解释为可交易资产；执行范围仍然是510300或现金。",
            "报告是研究证据汇总，不是当前交易信号、持仓建议或订单指令。",
        ],
    )


def build_document() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = "510300历史策略与全部回测结果报告"
    doc.core_properties.subject = "510300历史有效性、回测、机制与目标差距"
    doc.core_properties.author = "Codex研究整理"
    doc.core_properties.keywords = "510300, 历史回测, 夏普率, 微观结构, 估值, 机制图谱"

    add_cover(doc)
    add_navigation(doc)
    section_executive_summary(doc)
    section_methodology(doc)
    section_contract(doc)
    section_microstructure(doc)
    section_valuation(doc)
    section_registered_factors(doc)
    section_other_strategies(doc)
    section_intraday(doc)
    section_atlas(doc)
    section_final(doc)
    section_appendices(doc)

    # 防止标题孤立在页尾。
    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3", "Title", "Subtitle"}:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
    print(f"bytes={path.stat().st_size}")
    print(f"sha256={sha256(path)}")
