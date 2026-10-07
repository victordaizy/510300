"""汇总第一批保存结果、制度影响和任务落实，不重训或重跑账户。"""
from pathlib import Path
import json
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from research.daily_native_baseline_v1 import configs, minute_fields
from research import intraday_process_increment_v1 as old


def main():
    cfg, parent, out, prior = configs()
    frozen = json.loads((out / "freeze.json").read_text(encoding="utf-8"))
    for item in frozen["files"]:
        old.require(old.digest(Path(item["path"])) == item["sha256"], "冻结文件变化：" + item["path"])
    summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    review = json.loads((out / "review_reconciliation.json").read_text(encoding="utf-8"))
    metric = pd.read_csv(out / "11_D_common与native完整账户.csv")
    native = pd.read_parquet(out / "native_accounts.parquet")
    previous = pd.read_parquet(prior / "08_完整账户逐日账本.parquet")
    previous = previous.loc[previous.model.eq("D")].copy()
    paired = native.merge(previous, on=["date", "capital", "cost"], suffixes=("_native", "_common"), validate="one_to_one")
    paired["cumulative_native_minus_common_cny"] = paired.equity_native - paired.equity_common
    paired["daily_native_minus_common_cny"] = paired.groupby(["capital", "cost"]).cumulative_native_minus_common_cny.diff().fillna(paired.cumulative_native_minus_common_cny)
    paired[["date", "capital", "cost", "equity_common", "equity_native", "cumulative_native_minus_common_cny", "daily_native_minus_common_cny"]].to_csv(out / "17_native相对common人民币归因.csv", index=False, encoding="utf-8-sig")
    q = pd.read_parquet(out / "runtime_states.parquet")
    coverage = {"input_days": {key: int(q[key + "_input_valid"].sum()) for key in ("D", "A", "B", "C")}, "D_common_input_days": int(q.D_common_input_valid.sum()), "native_model_state_counts": q.model_state.value_counts().to_dict(), "native_prediction_state_counts": q.prediction_state.value_counts().to_dict(), "first_native_prediction": q.loc[q.prediction_D.notna(), "date"].min(), "execution_open_states": q.open_execution_data_state.value_counts().to_dict(), "execution_close_states": q.close_execution_data_state.value_counts().to_dict(), "A_B_C_native_models_run": False}
    old.write_json(out / "18_覆盖与资格汇总.json", coverage)
    old.write_json(out / "tests_receipt.json", {"command": "python -X utf8 -m pytest tests/test_daily_native_baseline_v1.py tests/test_intraday_process_increment_v1.py -q", "result": "12 passed in 5.00s", "performed_before_native_freeze_and_run": True, "new_tests": 7, "original_tests": 5, "development_fix": "冻结前首测发现pandas只读数组的原位布尔操作，复制数组后全部通过；未改变金融口径。"})
    old.write_json(out / "metric_contract.json", cfg["metric_contract"])
    rule_url = "https://www.sse.com.cn/aboutus/mediacenter/hotandd/c/c_20260424_10816474.shtml"
    matrix = [
        ("D日收益/日内收益/振幅/收盘位置", "日频定义可保持，收盘形成过程改变", "不机械重置504日，允许单一条件迁移", "可能分布漂移；28日独立描述，不作为验证"),
        ("D日金额及相对20日金额", "全日成交汇总定义保持；阶段权重可能变", "原单位与完整日检查，延续历史窗口", "不证明已辨认日内对齐误差"),
        ("D20日波动/动量", "日频统计公式不变，制度可影响分布", "保留20日公式，不新增参数", "新制度表现按原预定区段单列"),
        ("A五分钟冲击与30分钟修复", "最晚观察在旧协议14:45前，未进入收盘段", "本批只显示旧输入资格，保持原60日分制度预热", "未重建新A字段或训练；进一步迁移另立用途"),
        ("B相邻窗口金额推进", "窗口在收盘之前，量额误差仍需独立处理", "保留原金额门，禁止带动D输入失效", "未解决来源舍入/对齐；不重新运行B"),
        ("C固定尾盘30分钟", "收盘集合竞价直接改变尾段价格与量额过程", "原当日及历史60日资格保持", "不能继承旧尾段分布为已验证迁移"),
        ("次日开盘执行", "本次规则修订未据此宣称开盘字段发生相同变化", "原09:30价格/数量/量额检查，下一日才评估执行", "实际排队与个人成交未证明"),
        ("固定收盘退出", "从连续收盘阶段转为收盘集合竞价", "请求逐笔标注新旧机制，保持原条件价格/容量模拟", "15:00柱量不等于账户可成交保证"),
        ("盘后固定价格", "适用范围扩展至ETF", "本批不用此交易阶段", "原队列/估值/时钟资格仍受阻")
    ]
    pd.DataFrame(matrix, columns=["字段或用途", "变化影响", "本批处理", "边界"]).assign(官方来源=rule_url).to_csv(out / "19_制度影响矩阵.csv", index=False, encoding="utf-8-sig")
    old.write_json(out / "official_rule_reference.json", {"accessed_on": "2026-10-02", "source": rule_url, "rule_notice": "https://www.sse.com.cn/lawandrules/sselawsrules2025/stocks/exchange/c/c_20260424_10816482.shtml", "effective_date": "2026-07-06", "confirmed_scope": "基金收盘改为集合竞价、盘后固定价格范围扩展；不证明本地字段映射或真实成交", "matrix_migration_choices_are_research_assumptions": True})

    source = json.loads((old.ROOT / "reports/research/510300_pressure_recovery_v1/source_admission_20261002/summary.json").read_text(encoding="utf-8"))
    old.require(source["source_dates"] == 181 and source["source_files"] == 543, "盘口来源摘要范围变化")
    old.write_json(out / "orderbook_status_note.json", {"source_dates_recorded": source["source_dates"], "source_files_recorded": source["source_files"], "strict_M1_reference_days": source["strict_M1_reference_days"], "strict_M2_field_slots": source["strict_M2_field_slots"], "status": "SOURCE_RECORDS_EXIST_USE_SPECIFIC_EVIDENCE_MISSING", "full_raw_streams_rechecked_this_batch": False, "automatic_source_retry": False})
    amounts = minute_fields(pd.read_parquet(old.ROOT / parent["inputs"]["minutes"]))
    raw_distance = np.maximum(amounts.low * amounts.vol - amounts.amount, amounts.amount - amounts.high * amounts.vol).clip(lower=0)
    boundary = amounts.loc[amounts.strict_amount_bad & raw_distance.le(2 + 1e-8)].copy()
    boundary["outside_range_cny"] = raw_distance.loc[boundary.index]
    boundary.to_csv(out / "20_两元严格容差浮点边界.csv", index=False, encoding="utf-8-sig")
    old.write_json(out / "amount_boundary_explanation.json", {"strict_original_count": int(amounts.strict_amount_bad.sum()), "equivalent_amount_comparison_above_two_cny_with_float_epsilon": int(raw_distance.gt(2 + 1e-8).sum()), "floating_boundary_rows": len(boundary), "more_than_one_tick_bad_unchanged": int(amounts.amount_bad.sum()), "source_values_changed": False, "execution_or_prediction_changed": False, "interpretation": "原严格价比公式与等价金额边界在两个恰好约2元的样本有浮点比较差；仅记录，不修原42根一档异常或改变准入。"})

    workbook = load_workbook(out / "received/510300_48项后续工作清单.xlsx", read_only=True, data_only=False)
    tasks = []
    for row in workbook["48项任务"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and re.fullmatch(r"[GADRTPFM]\d{2}", row[0]):
            tasks.append(dict(zip(["任务ID", "工作包", "优先级", "任务", "输入", "交付物", "完成标准", "前置依赖", "外部原状态", "限制与说明", "原负责人"], row)))
    workbook.close()
    old.require(len(tasks) == 48, "收到的工作表不是48项")
    completed = {"A01": "review_reconciliation.json：原参数1820预测本地复算；旧账户保存核验已在前轮完成", "A02": "01_原108日原因表.csv", "A03": "availability_contract_v2.json与12_每日四层运行状态.csv", "A04": "封卷修正与设计分类.md", "A05": "02/03_C相对D人民币表", "A06": "orderbook_status_note.json：核来源摘要，不重读三流", "D02": "06_分钟字段质量分离.csv及依赖合同", "D04": "daily_asof_panel.parquet与标签成熟/未来不改前史测试；实际接收仍未知", "D05": "19_制度影响矩阵.csv；D条件迁移，其余不自动迁移", "D06": "一个D-native版本、36拟合、四完整账户，不晋升", "D07": "12项测试及本地status命令"}
    partial = {"G02": "本批主242/桥接252、rf/现金0；不覆盖其他研究线或实盘费率", "G03": "登记本研究家族的原失败/设计诊断；不是整个96因子库清理", "D01": "本批24项数据资产已分级；不是所有共享研究线的全量目录", "D03": "单位/精度/累计与日聚合完成有限诊断，逐行对齐及源精度原因未确认"}
    for row in tasks:
        key = row["任务ID"]
        row["本批状态"] = "本批完成" if key in completed else ("本批范围已落实，完整项仍有边界" if key in partial else "不在用户本批选择范围")
        row["本地依据或边界"] = completed.get(key, partial.get(key, "保留外部原状态，未启动"))
    pd.DataFrame(tasks).to_csv(out / "21_48项清单本批落实对照.csv", index=False, encoding="utf-8-sig")
    registry = [
        {"id": "510300_INTRADAY_PROCESS_INCREMENT_V1", "type": "FROZEN_ORIGINAL_COMPARISON", "status": "NO_CONFIRMED_INCREMENT_KEEP_FIXED_NEGATIVE_RESULTS", "changed": False},
        {"id": "OLD_CLOSE_PRESSURE_LOGISTIC", "type": "FROZEN_PRIOR", "status": "REJECTED_GOVERNANCE_GATES_FAILED", "changed": False},
        {"id": "OLD_RV20_MINUTE_SWITCH", "type": "FROZEN_PRIOR", "status": "REJECTED_FROZEN", "changed": False},
        {"id": cfg["study_id"], "type": cfg["classification"], "status": summary["status"], "configurations": 1, "new_fits": summary["native_model_fits"], "new_accounts": 4, "promotion_allowed": False},
    ]
    (out / "study_registry.jsonl").write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in registry) + "\n", encoding="utf-8")
    disposition = """# 封卷修正与设计分类

已接收外部Pro路线图/清单与粘贴复核文字，四项主要事实已用本地保存文件核实。原审阅ZIP与原16账户不改写。

|事项|分类|处置|
|---|---|---|
|108日D字段有效但共同门失效|原协议有意的共同比较资格，不是暗中实现错误|保留D-common，另立D-native|
|原共同日线不能代表独立日线运行|范围解释修正|报告分列共同对照和自有资格；训练成员与时钟改变明确登记|
|把总利润集中用于解释C边际收益|原文字解释不充分，需要更正|C−D总差3729.8021元，2024年9月为0；原增量未确认仍保持|
|当日标签必须成熟才能预测的说法|没有发现对应逻辑|训练需历史成熟标签；当日未来标签不作为预测资格，新增因果测试|
|盘口资料存在但资格受阻|来源状态澄清|181日/543文件有来源记录，严格用途未通过；不重复导出原数据|
|D-native改变输入、训练成员及制度迁移|新设计诊断|单一配置，新模型和四账户，不能称旧结果纠错后收益|
|严格金额公式两条浮点边界|数值表达诊断|保留原计数和数据，一档42条资格未变，不以此救结果|

本批没有A/B/C新策略、没有改变原模型超参数和执行政策。完整前瞻依赖实际收到时间与新日期，仍未建立。下一批未启动。
"""
    (out / "封卷修正与设计分类.md").write_text(disposition, encoding="utf-8")
    dependency = """# 依赖与运行合同

```mermaid
flowchart TD
  Daily[日线OHLC、金额与分红] --> DInput[D自身八项输入资格]
  MinutePrice[分钟价格与原过程预热] --> AInput[A输入资格]
  MinuteAmount[相邻窗口金额资格] --> BInput[B输入资格]
  TailAmount[当日与过去60日尾盘资格] --> CInput[C输入资格]
  DInput --> Common[原D-common共同比较门]
  AInput --> Common
  BInput --> Common
  CInput --> Common
  DInput --> Train[仅D合格且标签成熟的训练成员]
  Train --> Model[504成熟日、20日更新、原岭模型]
  Model --> Predict[D-native盘后预测]
  Predict --> Decision[按现金、持仓和原费用形成次日请求]
  Decision --> Execute[次日价格、量额一致性、容量、T+1检查]
  Execute --> Account[成交、未成交或部分成交及完整库存]
```

图中A/B/C只展示输入依赖，本批未训练其独立模型。D-native仍使用日线金额变量；独立于C尾盘金额不等于独立于所有金额。预测形成时，次日执行资格尚未知。原请求不能根据最终开盘价格事后撤销再声称按开盘买到。

运行状态入口：scripts/run_510300_daily_native_baseline_v1.py，使用`--stage status --date YYYY-MM-DD`。原点输入表和1211日四层状态均已保存。没有对应日期时返回LOCAL_DATE_NOT_COVERED/NO_VIEW_NO_CURRENT_LOCAL_RECEIPT，实际持仓UNKNOWN；不会把历史空仓状态当作当前仓位。该入口是本地历史与覆盖诊断，不是自动获取行情或券商服务。D08/F系列仍未实施。
"""
    (out / "依赖与运行合同.md").write_text(dependency, encoding="utf-8")
    drop = summary["paired_mse_reduction"] / summary["paired_mse_common"]
    lines = ["# 第一批完成：封卷修正、依赖拆分与D-native基准", "", "第一批研发已完成。D-native消除了旧共同门对独立日线预测的连带阻断，但完整账户收益变差；可运行和有优势是不同结论。未达到净夏普1.2和复合年化10%，没有晋升策略或启动下一批。", "", "## 四项审阅纠正", "", "本地用95组保存参数复算1820个原预测值，误差在数值精度内。108日确认为80日C当期/60日历史金额资格和28日新制度预热；其D八项字段均有限。没有发现当日必须已知自身未来收益才允许预测的逻辑。原首信号2024-08-29、首账户日2024-08-30解释364/363差。", "", "20万元BASE的C−D最终权益差为3729.8021元，2024年9月差额0。旧报告关于总利润集中度的数值保持，但不能由此解释C的边际差额；最大正月差发生在2025年11月，为2873.72338元。月差还包含资金路径后效，不是独立新机会或alpha证明。C的原区间仍跨零，未重新抽样或晋升。", "", "181日/543盘口文件的来源记录确实存在；本批只核对旧来源摘要，不重读全部三流。严格时钟、字段、估值及成交用途资格仍受阻。", "", "## 资格与实现结果", "", f"1211个原点中，D/A/B/C输入合格分别{coverage['input_days']['D']}/{coverage['input_days']['A']}/{coverage['input_days']['B']}/{coverage['input_days']['C']}日，原共同门869日。D-native累计706日可预测，第一日2023-09-11；前505日明确为训练不足。原评价首信号至末日472个预测、471个账户日全部有预测，原108日连带NO_VIEW消失。A/B/C只拆资格，不运行新模型。", "", "D特征从日线/分红直接重建，与原八项逐值一致；原训练起点2021-08-12和alpha1/504成熟日/20日更新/±5保持。新训练集与首次拟合时钟改变，36次拟合是一个版本的顺序更新，不是36个候选。原开收盘、费用、T+1、容量、固定到期、分红和完整日历合同复用。", "", "日频定义跨2026-07-06延续是一项条件迁移假设，不证明市场分布不变；尾盘及执行机制单独登记。详情见19_制度影响矩阵.csv。12项必要测试通过，4账户1884日账本和294个合并情景周期复算通过。", "", "## 同一471日完整账户", "", "主口径242交易日、现金/rf均0；252只做原净值桥接，不用于选优。费用为研究假设，2万元/20万元不是确认的实盘资金。", "", "|本金|费用|模型|年化收益|净夏普|最大回撤|周期|周期等权净期望|", "|---:|---|---|---:|---:|---:|---:|---:|"]
    for row in metric.sort_values(["capital", "cost", "model"]).itertuples():
        lines.append(f"|{row.capital:,}|{row.cost}|{row.model}|{row.annualized_return:.3%}|{row.net_sharpe:.3f}|{row.max_drawdown:.3%}|{row.completed_cycles}|{row.net_expectancy:.4%}|")
    base = metric.loc[metric.capital.eq(200000) & metric.cost.eq("BASE")].set_index("model")
    cost_delta = (base.loc["D_NATIVE", ["commission_cny", "slippage_cny"]].sum() - base.loc["D_COMMON", ["commission_cny", "slippage_cny"]].sum())
    gross_delta = base.loc["D_NATIVE", "completed_gross_quote_pnl"] - base.loc["D_COMMON", "completed_gross_quote_pnl"]
    lines += ["", f"原364个共同预测日期上，D-native的MSE从{summary['paired_mse_common']:.10f}降至{summary['paired_mse_native']:.10f}，点值改善{drop:.3%}；未重抽区间，不能称显著优势。完整账户明显变差，证明平均预测误差改善不能替代交易政策与费用后的价值验证。", "", f"20万元BASE的两条实际模拟路径相比，D-native毛报价损益少{abs(gross_delta):,.2f}元，佣金与滑点合计多{cost_delta:,.2f}元。净差因此同时包含交易选择/持有路径的毛收益下降与额外成本，不能只归咎手续费。原80日C缺口期间净损益−14928.86554元，新制度28日−11707.85834元；它们是事前登记的区段描述，不能据此恢复旧C门作避险信号。", "", "周期等权净均值、按日算术收益与全账户复合收益是不同统计。压力场景周期均值略正但全账户净值下降，不构成账户盈利。更多可用日期不保证更多有效机会；没有删亏损日、改费用、改期限或用最优年化常数救结果。", "", "## 数据限制及运行边界", "", "单位及全日量额没有发现整体100/1000倍错配，同一交易日相邻分钟金额有146915次下降，不支持把全部记录当连续累计值。源金额为整数，但源精度/逐字段对齐仍未知；不能据此把全部异常归因于取整。原严格计数5489与等价金额容差比较5487相差两条浮点边界，原一档42条异常资格未改。未知保留，不用量乘收盘补金额。", "", "当前分钟原点覆盖止于2026-08-12。运行2026-10-02状态命令返回NO_VIEW_NO_CURRENT_LOCAL_RECEIPT，预测为空、实际仓位未知；没有冒充当天观点。接收时钟未证实，前瞻采集与真实执行依然是之后的工作。", "", "## 本批范围与停止线", "", "21_48项清单本批落实对照.csv逐项保留外部原状态。A01—A06、D02/D04—D07完成本批明确工作；G02/G03/D01/D03注明本批口径、家族范围、24项资产和根因未识别边界。其余任务未启动，真实本金/风险和券商费用仍待实际实施阶段确认。本批固定诊断停止，不据结果调模型、窗口、退出或强行组合；下一批先做实质新信息及旧研究差异的准入，不立即运行新的收益实验。"]
    (out / "第一批研究结论与接续.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, axes = plt.subplots(2, 1, figsize=(11, 7), sharex=True)
    for axis, capital in zip(axes, [200000, 20000]):
        for data, label, color in [(previous, "原共同门 D-common", "#315e9b"), (native, "自有资格 D-native", "#d4712f")]:
            part = data.loc[data.capital.eq(capital) & data.cost.eq("BASE")]
            axis.plot(part.date, part.equity / capital, label=label, color=color, linewidth=1.6)
        axis.axvspan(pd.Timestamp("2026-03-09"), pd.Timestamp("2026-08-12"), color="#9c9c9c", alpha=.12, label="原共同门108日无预测")
        axis.axvline(pd.Timestamp("2026-07-06"), color="#777777", linestyle=":", linewidth=1)
        axis.set_title(f"{capital // 10000}万元 · 基础费用 · 完整471日")
        axis.set_ylabel("账户净值")
        axis.grid(alpha=.2)
        axis.legend(fontsize=8, loc="upper left")
    fig.suptitle("D-native覆盖增加，收益未改善：历史条件模拟", fontsize=14)
    fig.tight_layout()
    fig.savefig(out / "D_common与native净值.png", dpi=150)
    plt.close(fig)
    old.write_json(out / "completion_receipt.json", {"status": "COMPLETED_USER_SELECTED_FIRST_BATCH", "authority": cfg["user_scope"], "explicit_completed_task_ids": list(completed), "bounded_task_ids": list(partial), "other_tasks_started": False, "native_code_or_config_changed_after_results": False, "original_67_frozen_files_unchanged": True, "new_diagnostic_configurations": 1, "native_fits": 36, "native_accounts": 4, "new_tests_passed": 7, "original_tests_passed": 5, "financial_goal_achieved": False, "independent_validation": "NOT_ESTABLISHED", "actual_orders": 0, "new_market_downloads": 0})
    files = [{"path": str(p.relative_to(out)), "bytes": p.stat().st_size, "sha256": old.digest(p)} for p in sorted(out.rglob("*")) if p.is_file() and p.name != "file_index.json"]
    total = sum(x["bytes"] for x in files)
    old.require(total <= cfg["maximum_output_bytes"], "第一批报告超过存储预算")
    old.write_json(out / "file_index.json", {"files": files, "total_bytes": total, "index_excludes_itself": True})
    print(json.dumps({"状态": "第一批研究报告与任务对照已保存", "MSE点改善": drop, "文件数": len(files), "字节": total, "没有启动第二批": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
