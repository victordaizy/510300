"""复核四表与保存工作簿的实际值；不计算不存在的策略收益。"""

import csv
import hashlib
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/research/510300_pressure_recovery_v1"
WORKBOOK = ROOT / "outputs/01a0f739-6bb2-7741-8d30-204e31d95c62/510300_价格让步与压力恢复_四表.xlsx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(name):
    with (REPORT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def main():
    observations = read_csv("01_同步行情表.csv")
    events = read_csv("02_事件表.csv")
    orders = read_csv("03_订单表.csv")
    results = read_csv("04_结果表.csv")
    coverage = read_csv("每日覆盖与缺失.csv")
    status = json.loads((REPORT / "status.json").read_text(encoding="utf-8"))
    checks = {
        "raw_record_count": len(observations) == status["raw_observations"] == 957,
        "raw_distinct_dates": len({r["trade_date"] for r in observations}) == len(coverage) == 8,
        "coverage_rows_reconcile": sum(int(r["raw_rows"]) for r in coverage) == len(observations),
        "all_raw_observations_unqualified": all(r["quality_status"] == "NO_VIEW" for r in observations),
        "no_invented_event_or_order": len(events) == len(orders) == 0,
        "all_six_results_retained": len(results) == 6,
        "missing_results_not_zero": all(r[field] == "" for r in results for field in ("fill_rate", "net_expectancy", "win_rate", "payoff_ratio", "net_sharpe", "max_drawdown")),
        "missing_event_days_not_declared_no_event": status["known_no_event_days"] == 0 and status["unknown_event_days"] == 8,
    }
    registration = json.loads((REPORT / "registration_receipt.json").read_text(encoding="utf-8"))
    checks["config_identity_unchanged"] = sha(ROOT / "config/510300_pressure_recovery_v1.json") == registration["config_sha256"]
    checks["engine_identity_unchanged"] = sha(ROOT / "research/pressure_recovery_v1.py") == registration["engine_sha256"]
    snapshots = []
    for row in json.loads((REPORT / "source_inventory.json").read_text(encoding="utf-8")):
        if "sha256" not in row:
            continue
        source = ROOT / row["path"]
        saved = REPORT / "input_snapshot" / row["path"]
        snapshots.append(sha(source) == sha(saved) == row["sha256"])
    checks["source_snapshots_match"] = all(snapshots) and len(snapshots) == 3
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=False)
    checks["workbook_eight_sheets"] = wb.sheetnames == ["研究结论", "同步行情", "事件", "订单", "结果", "日期覆盖", "费用演示", "数据来源"]
    checks["workbook_957_records"] = sum(row[0] is not None for row in wb["同步行情"].iter_rows(min_row=5, values_only=True)) == 957
    checks["workbook_empty_events_orders"] = all(sum(row[0] is not None for row in wb[name].iter_rows(min_row=5, values_only=True)) == 0 for name in ("事件", "订单"))
    checks["cost_formula_cached_values"] = [wb["费用演示"][f"G{r}"].value for r in (5, 6, 7)] == [30, 15, 14]
    checks["costs_remain_formulas"] = all(formulas["费用演示"][f"G{r}"].data_type == "f" for r in (5, 6, 7))
    received = wb["同步行情"]["C5"].value
    original_received = datetime.fromisoformat(observations[0]["received_at"]).astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None)
    checks["received_time_preserves_china_clock"] = isinstance(received, datetime) and abs((received-original_received).total_seconds()) < .002
    checks["exit_request_is_0935"] = wb["日期覆盖"]["G5"].value == datetime(2026, 8, 13, 9, 35)
    errors = []
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    checks["no_saved_excel_error_cells"] = not errors
    wb.close()
    formulas.close()
    receipt = {
        "status": "PASS_SAVED_TABLES_AND_WORKBOOK" if all(checks.values()) else "FAILED",
        "checks": checks, "excel_errors": errors,
        "targeted_unit_tests_observed": {"command": ".venv\\Scripts\\python.exe -m pytest tests\\test_pressure_recovery_v1.py -q -p no:cacheprovider", "passed": 21, "failed": 0, "test_sha256": sha(ROOT / "tests/test_pressure_recovery_v1.py"), "evidence": "本任务此前实际运行输出，非本脚本重新运行"},
        "workbook": {"path": str(WORKBOOK.relative_to(ROOT)), "bytes": WORKBOOK.stat().st_size, "sha256": sha(WORKBOOK)},
        "artifact_process": {"exit_code_observed": 1, "export_message_printed": True, "saved_file_reopened_and_checked": True},
        "visual_review": {"workbook_sheets": 8, "md102_pdf_physical_pages": [139, 140]},
        "independent_strategy_validation": False, "real_trade_performance_validated": False,
    }
    (REPORT / "verification.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    records = [{"path": str(p.relative_to(REPORT)), "bytes": p.stat().st_size, "sha256": sha(p)}
               for p in sorted(REPORT.rglob("*")) if p.is_file() and p.name != "FILE_INDEX.csv"]
    with (REPORT / "FILE_INDEX.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(records)
    print(json.dumps({"status": receipt["status"], "checks": len(checks), "failed_checks": [k for k, v in checks.items() if not v]}, ensure_ascii=False))
    if not all(checks.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
