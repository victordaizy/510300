"""完整原生导出两次耗尽内存后，流式保存全部四表并提取保存文件的页面视图。"""

from __future__ import annotations

import copy
import gzip
import hashlib
import json
import os
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


ROOT = Path(__file__).resolve().parents[1]
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
ET.register_namespace("", M[1:-1])


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    cfg = json.loads((ROOT / "config/510300_pressure_four_table_delivery_v1.json").read_text(encoding="utf-8"))
    report, artifact = ROOT / cfg["report_directory"], ROOT / cfg["artifact_directory"]
    for name in ["authoring_attempt_01.json", "authoring_attempt_02.json"]:
        failure = json.loads((artifact / name).read_text(encoding="utf-8"))
        if failure["xlsx_successfully_exported"] or failure["exit_code"] != 1:
            raise SystemExit("原生完整导出不可用的回执不成立，不启用回退作者。")
    output = artifact / cfg["workbook_filename"]
    if output.exists():
        raise SystemExit("工作簿已存在，拒绝覆盖。")
    payload = report / "workbook_data.json.gz"
    with gzip.open(payload, "rt", encoding="utf-8") as handle:
        data = json.load(handle)
    wb = Workbook(write_only=True)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    font = Font(name="Microsoft YaHei", size=10, color="27364A")
    heading_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Microsoft YaHei", size=16, bold=True, color="172940")
    note_font = Font(name="Microsoft YaHei", size=10, italic=True, color="526276")
    fill = PatternFill(fill_type="solid", fgColor="263B55")
    header_border = Border(bottom=Side(style="thin", color="D8E2EC"))
    text_align = Alignment(vertical="center", wrap_text=True)
    numeric_align = Alignment(horizontal="right", vertical="center")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "同步行情": [14, 15, 24, 20, 14, 16, 62, 23, 21, 23, 19, 24, 23, 21, 25, 21, 25, 20, 27, 26, 28, 23, 23, 23, 26, 26, 25, 24, 80],
        "事件": [25, 20, 15, 24, 24, 30, 45, 45, 30, 26, 30],
        "订单": [24, 24, 20, 34, 12, 24, 20, 24, 24, 24, 24, 24, 24, 30, 60],
        "结果": [46, 24, 20, 20, 23, 23, 23, 18, 18, 18, 21, 21, 21, 23, 21, 26, 58],
        "日期覆盖": [15, 16, 23, 25, 25, 22, 28, 29, 29, 29, 24, 38, 26],
        "数据来源": [27, 20, 15, 20, 23, 38, 23, 80],
        "字段说明": [39, 94], "公开来源": [30, 94, 30, 62], "研究状态": [35, 20, 81]
    }
    names = ["研究状态", *data["tables"]]
    for name in names:
        sh = wb.create_sheet(name)
        sh.sheet_view.showGridLines = False
        for index, width in enumerate(widths[name], 1):
            sh.column_dimensions[get_column_letter(index)].width = width
        if name != "研究状态" and len(data["tables"][name]["rows"]) > 20:
            sh.freeze_panes = "D6" if name == "同步行情" else "A6"
    def put(sheet, values: list, row_number: int, kind: str = "body", formats: dict | None = None) -> None:
        sheet.row_dimensions[row_number].height = 43 if kind == "header" else 48 if sheet.title in {"同步行情", "数据来源"} else 37
        cells = []
        for index, value in enumerate(values):
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = title_font if kind == "title" else note_font if kind == "note" else heading_font if kind == "header" else font
            if kind == "header":
                cell.fill, cell.alignment, cell.border = fill, header_align, header_border
            else:
                cell.alignment = numeric_align if isinstance(value, (int, float)) and not isinstance(value, bool) else text_align
            if formats and index in formats:
                cell.number_format = formats[index]
            cells.append(cell)
        if kind in {"title", "note"}:
            sheet.row_dimensions[row_number].height = 30 if kind == "title" else 25
            # 标题/说明在同一行后续空单元格自然显示，避免把独立字段合并。
            cells[0].alignment = Alignment(vertical="center", wrap_text=False)
        sheet.append(cells)
    front = wb["研究状态"]
    put(front, ["510300 压力修复研究四表"], 1, "title")
    put(front, ["2026-01-05至09-30；181日最新来源覆盖，原M1/M2尚未准入。"], 2, "note")
    front.append([])
    front.append([])
    put(front, ["项目", "当前记录", "含义"], 5, "header")
    rows = [
        ["来源日期数", "=COUNTA('日期覆盖'!$A$6:$A$186)", "全部181日逐日覆盖；不是181个合格事件日。"],
        ["三流来源文件数", "=COUNTA('数据来源'!$A$6:$A$548)", "行情、逐笔委托、逐笔成交各181文件。"],
        ["三流原始记录数", "=SUM('数据来源'!$G$6:$G$548)", "原件清单行数之和，完整源文件身份见数据来源。"],
        ["测量网格数", "=COUNTA('同步行情'!$C$6:$C$47789)", "全部264格/日，含连续分钟及收盘、盘后覆盖探针。"],
        ["已准入原时点的网格数", '=COUNTIF(\'同步行情\'!$F$6:$F$47789,"QUALIFIED")', "0表示原字段/时点门槛未通过，不表示市场没有机会。"],
        ["M1合格参考日期数", 0, "正IOPV字段仅4日，单位/经济时点/历史可得性未验证。"],
        ["正式事件数", None, "检验未运行；空值未知，不按0次市场事件解释。"],
        ["已提交订单数", 0, "研究没有提交订单；订单表保留空表。"],
        ["已确认成交数", 0, "没有真实成交或独立验证重放证据。"],
        ["成交后净期望", None, "尚未计算；结果表保留3版本×2账户规模。"],
        ["20万元全账户夏普", None, "完整现金日、持仓与执行账本尚未运行。"],
        ["原研究目标", "未完成", "来源表与工作簿完成没有被记为策略有效。"],
        ["本次费用（美元）", 0, "使用已存资格矩阵，未发起新数据下载。"],
    ]
    for row_number, row in enumerate(rows, 6):
        put(front, row, row_number, formats={1: "0.00%" if row_number == 15 else "0.000" if row_number == 16 else "#,##0"})
    front.append([])
    put(front, ["缺失值、名义报价与正式事件的解释见“字段说明”。"], 20, "note")
    put(front, ["下一步：新增时点、估值或执行证据后，先验证小样本，再启动原固定实验。"], 21, "note")
    for name, descriptor in data["tables"].items():
        sh = wb[name]
        put(sh, ["510300 源行报价与时点资格" if name == "同步行情" else f"510300 {name}"], 1, "title")
        put(sh, [descriptor["note"]], 2, "note")
        sh.append([])
        sh.append([])
        put(sh, descriptor["headers"], 5, "header")
        if descriptor["rows"]:
            sh.auto_filter.ref = f"A5:{get_column_letter(len(descriptor['headers']))}{len(descriptor['rows'])+5}"
        formats = {index: "yyyy-mm-dd" for index in descriptor["date_columns"]}
        formats.update({index: "hh:mm" for index in descriptor["time_columns"]})
        if name == "同步行情":
            formats.update({8: "0", 9: "0", 10: "0", 11: "0", 12: "0.0000", 13: "0.000", 14: "#,##0.00", 15: "#,##0", 16: "0.000", 19: "0", 20: "0"})
        elif name == "结果":
            formats.update({index: "#,##0" for index in range(1, 7)})
            formats.update({7: "0.00%", 8: "0.00%", 9: "0.00%", 10: "0.000", 11: "0.00%", 13: "0.000", 14: "0.00%"})
        elif name == "日期覆盖":
            formats.update({index: "#,##0" for index in range(2, 10)})
        elif name == "数据来源":
            formats[6] = "#,##0"
        for number, row in enumerate(descriptor["rows"], 6):
            values = [datetime.fromisoformat(value).date() if value is not None and index in descriptor["date_columns"] else value
                      for index, value in enumerate(row)]
            put(sh, values, number, formats=formats)
        print(f"流式写入{name}：{len(descriptor['rows']):,}行。", flush=True)
    temporary = output.with_suffix(".streaming.tmp.xlsx")
    wb.save(temporary)
    cached = {"B6": len(data["tables"]["日期覆盖"]["rows"]), "B7": len(data["tables"]["数据来源"]["rows"]),
              "B8": sum(x[6] for x in data["tables"]["数据来源"]["rows"]), "B9": len(data["tables"]["同步行情"]["rows"]),
              "B10": sum(x[5] == "QUALIFIED" for x in data["tables"]["同步行情"]["rows"])}
    with zipfile.ZipFile(temporary) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                tree = ET.fromstring(raw)
                for cell in tree.iter(f"{M}c"):
                    if cell.get("r") in cached:
                        value = cell.find(f"{M}v")
                        if value is None:
                            value = ET.SubElement(cell, f"{M}v")
                        value.text = str(cached[cell.get("r")])
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            target.writestr(item, raw)
    temporary.unlink()
    # 页面摘录只用于检查最终保存文件的样式，完整工作簿保持全部行列。
    preview = artifact / "saved_workbook_viewports.xlsx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(preview, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                kept = []
                with source.open(item.filename) as handle:
                    parser = ET.iterparse(handle, events=["end"])
                    for _, element in parser:
                        if element.tag == f"{M}row":
                            if int(element.get("r")) <= 22:
                                row = copy.deepcopy(element)
                                for cell in row.findall(f"{M}c"):
                                    formula = cell.find(f"{M}f")
                                    if formula is not None:
                                        cell.remove(formula)
                                kept.append(row)
                            element.clear()
                    tree = parser.root
                sheet_data = tree.find(f"{M}sheetData")
                position = list(tree).index(sheet_data)
                tree.remove(sheet_data)
                replacement = ET.Element(f"{M}sheetData")
                replacement.extend(kept)
                tree.insert(position, replacement)
                last_col = get_column_letter(3 if item.filename == "xl/worksheets/sheet1.xml" else len(data["tables"][names[int(Path(item.filename).stem[5:])-1]]["headers"]))
                for tag in ["dimension", "autoFilter"]:
                    node = tree.find(f"{M}{tag}")
                    if node is not None:
                        node.set("ref", f"A{5 if tag == 'autoFilter' else 1}:{last_col}22")
                target.writestr(item, ET.tostring(tree, encoding="utf-8", xml_declaration=True))
            else:
                target.writestr(item, source.read(item.filename))
    result = {"saved_at": datetime.now().astimezone().isoformat(), "workbook": str(output), "bytes": output.stat().st_size,
              "sha256": sha(output), "compressed_data_sha256": sha(payload), "streaming_writer_sha256": sha(Path(__file__)),
              "builder_sha256": sha(artifact / "build_pressure_recovery_181day_workbook.mjs"), "sheet_names": names,
              "source_summary_formula_values": list(cached.values()), "authoring_method": "STREAMING_AFTER_TWO_NATIVE_EXPORT_MEMORY_FAILURES",
              "all_grid_rows_retained": 47784, "formal_event_count": None, "strategy_returns": "NOT_COMPUTED", "goal_achieved": False,
              "new_market_data_requests": 0, "fee_usd": 0, "preview_is_only_saved_file_viewport": True,
              "preview_formulas_replaced_by_saved_cache_for_visual_check_only": True}
    (artifact / "workbook_export_receipt.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"完整工作簿已流式保存：{output.stat().st_size:,}字节；5个覆盖公式缓存由来源逐项计算，收益仍未计算。", flush=True)


if __name__ == "__main__":
    main()
